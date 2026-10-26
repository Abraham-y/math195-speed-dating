"""Abraham's implementation slides for the restructured midterm deck: two after "BIC and Stepwise Logistic
Regression" (the full logistic model, then stepwise selection and the lasso), two after "LDA and QDA" (the two models,
then QDA's shrinkage) and one after the ROC slide (how the event-clustered error bars are computed). They are built by
cloning the styling (cards, headings, big-number card, result strip, footer, background) of the deck's own slides
"Hold out whole events" and "Too many predictors", so they match when imported into Google Slides (File > Import slides).

Input : the team deck downloaded as .pptx (--template; only its styling is used, none of its slides are kept);
        report/figures/slide_spline_bic.png, slide_stepwise_kept.png, slide_qda_shrink.png, slide_toy_cluster.png,
        slide_eq_bic.png, slide_eq_shrink.png, slide_eq_se.png (scripts/report_figures.py --slides)
Output: results/report_impl_slides.pptx: the five slides with speaker notes;
        with --full, results/report_impl_slides_full.pptx: the whole deck with the new slides in place of its two empty
        slides and after the ROC slide, the tuning slide moved into Results, four slide errors fixed (main result
        0.54 -> 0.58, QDA covariance "2775" -> "about 8,000" numbers, the stock ROC image -> our curves, the
        pre-registration footer) and footers renumbered; --skip marks named slides as skipped (hidden when presenting,
        not deleted) and leaves them out of the numbering. The team's speaker notes are kept exactly as they were.
Numbers on the slides: AUCs from results/part1_models_metrics.csv (headline, event 5-fold CV); column counts, the QDA
shrinkage chosen, lasso columns kept and the training sizes from results/part1_models_folds.csv; predictors kept from
results/part1_models_stepwise_selection.csv; spline fit gains from report_figures.spline_bic. Nothing is random.
Run from the repository root:
  python scripts/report_impl_slides.py --template "path/to/Speed-dating midterm presentation.pptx" [--full] [--skip TEXT ...]
"""
import argparse, copy, io, pathlib, re
from pptx import Presentation
from pptx.util import Inches
from pptx.enum.text import MSO_ANCHOR
from pptx.oxml.ns import qn

RES, FIG = pathlib.Path("results"), pathlib.Path("report/figures")


def find(slide, start):
    for sh in slide.shapes:
        if sh.has_text_frame and sh.text_frame.text.startswith(start):
            return sh
    raise KeyError(start)


def clone(shape, slide, left, top, width, height):
    el = copy.deepcopy(shape._element)
    tree = slide.shapes._spTree
    ext = tree.find(qn("p:extLst"))
    ext.addprevious(el) if ext is not None else tree.append(el)
    sh = slide.shapes[-1] if ext is None else next(s for s in slide.shapes if s._element is el)
    sh.left, sh.top, sh.width, sh.height = Inches(left), Inches(top), Inches(width), Inches(height)
    return sh


def set_text(sh, paras):
    """Replace the text, keeping the first paragraph's and first run's formatting. paras: strings or lists of (text, bold)."""
    body = sh.text_frame._txBody
    ps = body.findall(qn("a:p")); proto = ps[0]
    for p in ps:
        body.remove(p)
    for para in paras:
        runs = [(para, None)] if isinstance(para, str) else para
        p = copy.deepcopy(proto); rs = p.findall(qn("a:r")); rproto = rs[0]
        for e in rs + p.findall(qn("a:br")) + p.findall(qn("a:fld")):
            p.remove(e)
        end = p.find(qn("a:endParaRPr"))
        for text, bold in runs:
            r = copy.deepcopy(rproto); r.find(qn("a:t")).text = text
            if bold is not None:
                rpr = r.find(qn("a:rPr"))
                if rpr is None:
                    rpr = r.makeelement(qn("a:rPr"), {}); r.insert(0, rpr)
                rpr.set("b", "1" if bold else "0")
            end.addprevious(r) if end is not None else p.append(r)
        body.append(p)


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = " ".join(text.split())


def new_slide(prs, tmpl):
    s = prs.slides.add_slide(tmpl.slide_layout)
    bg = tmpl._element.cSld.find(qn("p:bg"))
    if bg is not None:
        s._element.cSld.insert(0, copy.deepcopy(bg))
    return s


def slide_with(prs, text):
    for s in prs.slides:
        if any(sh.has_text_frame and text in sh.text_frame.text for sh in s.shapes):
            return s
    raise KeyError(text)


def build(prs):
    """Append the implementation slides to prs, styled from its own slides. Returns {"logistic": [..], "lda": [..],
    "errors": [..]}: the slides that go in the two empty slots and the one that follows the ROC slide."""
    t19, t23 = slide_with(prs, "Hold out whole events"), slide_with(prs, "Too many predictors")
    eyebrow, title, footer = find(t19, "4 · TECHNICAL"), find(t19, "Hold out whole"), find(t19, "Math 195")
    card = next(sh for sh in t19.shapes if sh.has_text_frame and sh.text_frame.text == "" and sh.top < Inches(2.2))
    h3, body = find(t19, "Why"), find(t19, "People date up to")
    strip = next(sh for sh in t23.shapes if sh.has_text_frame and sh.text_frame.text == "" and sh.top > Inches(4.9))
    strip_text = find(t23, "Simpler beat fuller")
    dark = next(sh for sh in t23.shapes if sh.has_text_frame and sh.text_frame.text == "" and sh.height > Inches(2.5))
    big, big_caption = find(t23, "≈ 4"), find(t23, "matches per column")
    icon = next(sh for sh in t23.shapes if sh.shape_type == 13)
    icon_blob = icon.image.blob

    def slide(title_text, foot):
        s = new_slide(prs, t19)
        set_text(clone(eyebrow, s, 0.89, 0.89, 11.9, 0.26), ["4 · TECHNICAL IMPLEMENTATION"])
        set_text(clone(title, s, 0.89, 1.29, 11.9, 0.52), [title_text])
        set_text(clone(footer, s, 0.89, 6.88, 11.9, 0.26), [f"Math 195 · Speed dating · 0 · {foot}"])   # page set by assemble()
        return s

    def card_at(s, left, top, width, height, heading, text, text_h):
        clone(card, s, left, top, width, height)
        set_text(clone(h3, s, left + 0.2, top + 0.17, width - 0.4, 0.29), [heading])
        if text:
            set_text(clone(body, s, left + 0.2, top + 0.5, width - 0.4, text_h), [text])

    def strip_at(s, runs, left=0.9, top=5.8, width=11.54, height=0.88):
        clone(strip, s, left, top, width, height)
        s.shapes.add_picture(io.BytesIO(icon_blob), Inches(left + 0.22), Inches(top + height / 2 - 0.16), height=Inches(0.33))
        t = clone(strip_text, s, left + 0.62, top + 0.06, width - 0.75, height - 0.12); set_text(t, [runs])
        t.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE

    def big_number(s, left, top, width, height, number, caption):
        clone(dark, s, left, top, width, height)
        set_text(clone(big, s, left + 0.28, top + 0.3, width - 0.5, 1.0), [number])
        set_text(clone(big_caption, s, left + 0.28, top + 1.45, width - 0.5, height - 1.7), [caption])

    pic = lambda s, name, left, top, **size: s.shapes.add_picture(str(FIG / f"slide_{name}.png"), Inches(left), Inches(top),
                                                                    **{k: Inches(v) for k, v in size.items()})

    # ---- after "BIC and Stepwise Logistic Regression": the full model, then selection
    a1 = s = slide("Logistic regression on our data",
                   "BIC: James et al. (2023), ISL §6.1.3 · Splines: ISL §7.4 · Dots: the five pre-date training folds; no curve was chosen with the date ratings either")
    card_at(s, 0.9, 1.95, 4.95, 1.7, "What goes in",
            "Her answers, his answers and the pair features, such as the age gap: 74 pre-date predictors. Categories "
            "become yes/no columns, so about 125 columns in all.", 1.1)
    card_at(s, 0.9, 3.8, 4.95, 1.82, "Straight line or curve?",
            "For six predictors we also tried a curve (a spline), kept only if it lowered BIC:", 0.6)
    pic(s, "eq_bic", 1.1, 5.0, width=3.5)
    pic(s, "spline_bic", 6.15, 1.95, width=6.3)
    strip_at(s, [("A straight line won every time. ", True),
                 ("Full logistic regression, held-out AUC: 0.542 before the date, 0.780 with the date ratings.", False)])
    notes(s, """Jenna introduced logistic regression and BIC; here's how we ran them on our data. Each row is one pair.
        The predictors are her answers, his answers, and the pair features from earlier, such as the age gap and
        interest similarity: 74 in all before the date. Categories like field of study become yes-or-no dummy columns,
        so the model has about 125 columns. For six of the numeric predictors a straight-line effect might be too
        simple, so we also tried a curve, called a spline, for each one, and let BIC decide whether it stays. A curve
        adds two coefficients, and BIC charges log n for each, about 8 with roughly 3,300 training pairs, so about 16
        in all. The chart shows how much each curve actually improved the fit in each of the five training folds. His
        age came closest, at about 10, but no curve ever paid its price, so the model stayed linear everywhere. This
        full model reached a held-out AUC of 0.542 before the date and 0.780 with the date ratings.""")

    a2 = s = slide("Stepwise selection and the lasso on our data",
                   "Stepwise: ISL §6.1.2 · Lasso: ISL §6.2.2; Friedman et al. (2010) · Bars: pre-date fits · AUCs on held-out events")
    card_at(s, 0.9, 1.95, 4.95, 1.8, "Stepwise selection by BIC",
            "Add or drop one predictor at a time while BIC goes down, three ways (forward, backward, both). Kept about "
            "8 of 74 predictors per fit; about 15 of 90 with the date ratings.", 1.2)
    card_at(s, 0.9, 3.9, 4.95, 1.72, "The lasso",
            "Shrinks every coefficient toward zero and sets the weakest exactly to zero; inner CV picks how strongly. "
            "Kept 22 to 76 of about 125 columns.", 1.1)
    pic(s, "stepwise_kept", 6.15, 1.95, width=6.1)
    strip_at(s, [("Fewer predictors predicted better. ", True),
                 ("Before the date: full model 0.542, stepwise 0.559, lasso 0.573. With the date ratings: 0.780, 0.835, 0.838.", False)])
    notes(s, """Next we asked whether fewer predictors would do better. Stepwise selection adds or drops one predictor at
        a time, keeping a change only if it lowers BIC. We ran it three ways, forward, backward and both, in each of
        the five folds, so 15 fits, and it kept only about 8 of the 74 predictors each time. The chart shows the ones
        it kept most often. His self-rated fun, the age gap, how often he dates, sharing a field of study, and how much
        he values intelligence came up in most fits. Blue means more of it goes with more matches; orange, with fewer.
        We also ran the lasso, which you'll see in our results. It's logistic regression with a penalty that shrinks
        every coefficient toward zero and sets the weakest ones exactly to zero, and an inner cross-validation picks
        how strong the penalty is. It kept between 22 and 76 of the roughly 125 columns. Both beat the full model on
        new events: before the date, 0.559 for stepwise and 0.573 for the lasso, against 0.542; with the date ratings,
        0.835 and 0.838, against 0.780.""")

    # ---- after "LDA and QDA": the two models, then QDA's shrinkage
    b1 = s = slide("LDA and QDA on our data",
                   "LDA and QDA: James et al. (2023), ISL §4.4 · Counts from our training folds (event 5-fold CV): 124 to 128 columns, 3,047 to 3,750 pairs")
    card_at(s, 0.9, 1.95, 4.95, 1.7, "LDA",
            "One covariance matrix, shared by matches and non-matches, estimated from all of the about 3,300 training "
            "pairs. Nothing to tune.", 1.1)
    card_at(s, 0.9, 3.8, 4.95, 1.82, "QDA",
            "A separate covariance matrix for each class, so the matches' matrix has to come from the matches alone.", 1.1)
    big_number(s, 6.15, 1.95, 3.05, 3.67, "8,000", "numbers, roughly, in the matches' covariance matrix: 125 × 126 / 2 for 125 columns")
    big_number(s, 9.39, 1.95, 3.05, 3.67, "550", "matches, on average, in a training fold to estimate them from")
    strip_at(s, [("Too many numbers, too few matches: ", True),
                 ("plain QDA's matrix is mostly noise. LDA's held-out AUC: 0.553 before the date, 0.790 with the date ratings.", False)])
    notes(s, """Now LDA and QDA. Both describe each class, match and no match, by its average and its covariance matrix,
        the table of how every pair of predictors varies together. LDA uses one covariance matrix for both classes,
        estimated from all of the roughly 3,300 training pairs, with the training match rate, about 0.165, as the
        prior. It has nothing to tune. QDA gives each class its own covariance matrix, so the matches' matrix has to be
        estimated from the matches alone, and that's the problem. With about 125 columns the matrix has about 8,000
        separate numbers, 125 times 126 over 2, and a training fold has only about 550 matches to estimate them from.
        So most of those numbers are noise. LDA reached a held-out AUC of 0.553 before the date and 0.790 with the date
        ratings.""")

    b2 = s = slide("Shrinking QDA's covariance matrix",
                   "Regularized QDA: Friedman (1989); Hastie et al. (2009) §4.3.1 · Heatmaps: 8 of our predictors, 690 matched pairs")
    card_at(s, 0.9, 1.95, 4.95, 3.67, "Shrink toward no correlation",
            "Average the estimate with the identity matrix (no correlations), using weight r:", 0.6)
    pic(s, "eq_shrink", 1.1, 3.2, height=0.42)
    set_text(clone(body, s, 1.1, 3.85, 4.55, 1.6), ["r = 0 is plain QDA; r = 1 drops every correlation. Inner CV chose r "
                                                    "from 0.1, 0.3, 0.5, 0.7, 0.9, and picked 0.9 in all 10 fits."])
    pic(s, "qda_shrink", 6.05, 1.95, width=6.4)
    strip_at(s, [("Shrinking paid off. ", True),
                 ("Before the date: LDA 0.553, regularized QDA 0.579. With the date ratings: 0.790 and 0.836.", False)])
    notes(s, """Our fix is called regularized QDA. We shrink the matches' covariance matrix toward the identity matrix,
        which means no correlations at all, using a weighted average with weight r. With r equal to 0 it's plain QDA;
        with r equal to 1 every correlation is dropped. An inner cross-validation inside the training events chose r
        from 0.1, 0.3, 0.5, 0.7 and 0.9. The heatmaps show what shrinking does, for 8 of our predictors among the
        matched pairs. On the left are the correlations we estimate; for example, women who rate themselves as more
        attractive also rate themselves as more fun, a correlation of 0.44. On the right is what QDA uses after
        shrinking with r equal to 0.9: every correlation is about a tenth as big, so 0.44 becomes 0.04. The inner
        cross-validation picked 0.9, the strongest option, in all ten fits. And regularized QDA ended up one of our
        best models: 0.579 before the date against 0.553 for LDA, and 0.836 against 0.790 with the date ratings.""")

    # ---- after the ROC slide: how the error bars on our results are computed
    c = s = slide("Error bars that respect events",
                  "Cluster-robust SE: Cameron & Miller (2015) · Clustered AUC: Obuchowski (1997) · The toy numbers are made up")
    card_at(s, 0.9, 1.95, 5.9, 3.67, "Toy example: six pairs at two events", "", 0)
    pic(s, "toy_cluster", 1.05, 2.45, width=5.6)
    card_at(s, 7.0, 1.95, 5.44, 2.3, "Add up within each event, then square",
            "Pairs at an event share people and an evening, so errors lean the same way. ℓᵢ is pair i's error; "
            "G = 21 events, n = 4,184 pairs:", 0.95)
    pic(s, "eq_se", 7.2, 3.42, height=0.64)
    card_at(s, 7.0, 4.4, 5.44, 1.22, "On our results",
            "90% intervals are AUC ± 1.72 × SE (t with 20 df). The AUC's SE uses the same idea.", 0.6)
    strip_at(s, [("Ignoring events makes error bars too narrow ", True),
                 ("whenever errors lean together within an event.", False)])
    notes(s, """Every result we show comes with an error bar, so here's how we compute it. The usual standard error
        assumes all 4,184 pairs are independent. Ours aren't: pairs at the same event share people and the same
        evening, so a model's errors tend to lean the same way within an event. Here's a toy example with six pairs at
        two events. At event A the model is too high for every pair; at event B it's too low for every pair. The usual
        formula squares each pair's error and adds them up, which gives 0.28. The clustered version first adds up each
        event's errors, then squares, which gives 0.72. Errors that lean together don't cancel within an event, so the
        variance is more than twice as big and the error bar about 1.6 times as wide. That's the formula on the right:
        sum within each event, square, add over our 21 events, with a small correction of G over G minus 1. Our
        intervals are the AUC plus or minus 1.72 standard errors, from a t distribution with 20 degrees of freedom,
        and the AUC's standard error uses the same idea, due to Obuchowski. Whenever errors lean together within an
        event, ignoring the events makes the error bars too narrow.""")

    return {"logistic": [a1, a2], "lda": [b1, b2], "errors": [c]}


def replace_in_frame(tf, old, new):
    """Replace text in a text frame; a match split across runs is collapsed into the paragraph's first run."""
    hit = 0
    for p in tf.paragraphs:
        full = "".join(r.text for r in p.runs)
        if old not in full:
            continue
        one = next((r for r in p.runs if old in r.text), None)
        if one is not None:
            one.text = one.text.replace(old, new)
        else:
            p.runs[0].text = full.replace(old, new)
            for r in p.runs[1:]:
                r._r.getparent().remove(r._r)
        hit += 1
    return hit


def replace_text(slide, old, new):
    hit = 0
    for sh in slide.shapes:
        if sh.has_text_frame:
            hit += replace_in_frame(sh.text_frame, old, new)
        if sh.has_table:
            hit += sum(replace_in_frame(c.text_frame, old, new) for c in sh.table.iter_cells())
    assert hit, f"{old!r} not found"


def replace_image(slide, path):
    """Point the slide's (only) picture at a new image, keeping its height and centre."""
    from PIL import Image
    pic = next(sh for sh in slide.shapes if sh.shape_type == 13)
    _, rid = slide.part.get_or_add_image_part(str(path))
    fill = pic._element.blipFill
    fill.find(qn("a:blip")).set(qn("r:embed"), rid)
    crop = fill.find(qn("a:srcRect"))
    if crop is not None:
        fill.remove(crop)
    w, h = Image.open(path).size; cx = pic.left + pic.width // 2
    pic.width = int(pic.height * w / h); pic.left = int(cx - pic.width // 2)


def assemble(prs, new, skip=()):
    """The whole deck: the new slides replace the two empty ones (new["logistic"] the first, new["lda"] the second),
    new["errors"] follows the ROC slide, the tuning slide moves into Results after "Every model is close to chance",
    four slide errors are fixed, slides whose text contains any string in `skip` are marked as skipped (hidden when
    presenting, not deleted), and the footers of the shown slides are renumbered. Speaker notes of the team's slides
    are left exactly as they were."""
    ids = prs.slides._sldIdLst; slides = list(prs.slides); el = {int(e.get("id")): e for e in ids}
    new_ids = {s.slide_id for group in new.values() for s in group}
    blanks = [s for s in slides if s.slide_id not in new_ids and not any(sh.has_text_frame and sh.text_frame.text.strip() or sh.shape_type == 13 for sh in s.shapes)]
    assert len(blanks) == 2, f"expected two empty slides, found {len(blanks)}"
    swap = {blanks[0].slide_id: new["logistic"], blanks[1].slide_id: new["lda"]}
    order = [x for s in slides if s.slide_id not in new_ids for x in swap.get(s.slide_id, [s])]
    pos = lambda sl: next(i for i, s in enumerate(order) if s.slide_id == sl.slide_id)
    tuning, chance = slide_with(prs, "Too many predictors"), slide_with(prs, "Every model is close to chance")
    t = order.pop(pos(tuning)); order.insert(pos(chance) + 1, t)
    roc = slide_with(prs, "Model evaluation: the ROC curve")
    for j, x in enumerate(new["errors"]):
        order.insert(pos(roc) + 1 + j, x)
    blank_ids, order_ids = [s.slide_id for s in blanks], [s.slide_id for s in order]     # before any change
    for i in blank_ids:
        prs.part.drop_rel(el[i].get(qn("r:id")))
    for e in list(ids):
        ids.remove(e)
    for i in order_ids:
        ids.append(el[i])
    replace_text(tuning, "4 · WHAT THE TUNING TOLD US", "5 · RESULTS")
    main_result = slide_with(prs, "THE MAIN RESULT")
    big = [sh for sh in main_result.shapes if sh.has_text_frame and sh.text_frame.text.strip() == "0.54"]
    assert len(big) == 1; replace_in_frame(big[0].text_frame, "0.54", "0.58")      # best pre-date AUC is 0.583
    replace_text(slide_with(prs, "Methods we used beyond"), "(2775 numbers)", "(about 8,000 numbers)")   # ~125 columns
    replace_text(slide_with(prs, "The same pre-processing pipeline"), " · Pre-registration: Nosek et al. (2018), PNAS", "")
    hidden = {slide_with(prs, t).slide_id for t in skip}
    for s in order:
        if s.slide_id in hidden:
            s._element.set("show", "0")
    shown = [s for s in order if s.slide_id not in hidden]
    page = {s.slide_id: i for i, s in enumerate(shown, 1)}
    cv_slide = slide_with(prs, "Hold out whole events")
    replace_text(slide_with(prs, "Our folds are whole events"), "slide 17", f"slide {page[cv_slide.slide_id]}")
    replace_image(slide_with(prs, "Model evaluation: the ROC curve"), FIG / "slide_roc.png")       # our curves, not a stock image
    replace_image(cv_slide, FIG / "slide_cv_design.png")                                           # training set / test set
    for s in shown:
        i = page[s.slide_id]
        for sh in s.shapes:
            m = sh.has_text_frame and re.match(r"Math 195 · Speed dating · (\d+)(?=\s|$)", sh.text_frame.text)
            if m:
                replace_in_frame(sh.text_frame, m.group(0), f"Math 195 · Speed dating · {i}")
    return order


def footer_page(slide):
    return next(int(m.group(1)) for sh in slide.shapes if sh.has_text_frame
                for m in [re.match(r"Math 195 · Speed dating · (\d+)(?=\s|$)", sh.text_frame.text)] if m)


def main():
    ap = argparse.ArgumentParser(); ap.add_argument("--template", required=True)
    ap.add_argument("--full", action="store_true", help="write the whole deck with the new slides in place")
    ap.add_argument("--skip", nargs="*", default=[], help="with --full: mark slides containing these strings as skipped")
    args = ap.parse_args()
    prs = Presentation(args.template); n_old = len(prs.slides)
    new = build(prs)
    if args.full:
        order = assemble(prs, new, args.skip); out = RES / "report_impl_slides_full.pptx"
    else:
        full = Presentation(args.template); fnew = build(full); assemble(full, fnew, args.skip)
        page = {a.slide_id: footer_page(b) for k in new for a, b in zip(new[k], fnew[k])}
        for group in new.values():
            for sl in group:
                for sh in sl.shapes:
                    if sh.has_text_frame and sh.text_frame.text.startswith("Math 195 · Speed dating · 0 "):
                        replace_in_frame(sh.text_frame, "Math 195 · Speed dating · 0 ", f"Math 195 · Speed dating · {page[sl.slide_id]} ")
        ids = prs.slides._sldIdLst
        for sid in list(ids)[:n_old]:
            prs.part.drop_rel(sid.get(qn("r:id"))); ids.remove(sid)
        out = RES / "report_impl_slides.pptx"
    prs.save(out); print(f"wrote {out} ({len(prs.slides)} slides)")


if __name__ == "__main__":
    main()
