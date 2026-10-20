"""Abraham's slides for the midterm presentation (Technical Implementation and Next Steps), built inside the team
deck's own theme so they can be imported into Google Slides (File > Import slides) without restyling.

Input : the team deck downloaded as .pptx (--template; used only for its theme and layouts, none of its slides are kept),
        report/figures/slide_cv_design.png (scripts/report_figures.py --slides)
Output: results/report_slides.pptx: four slides with speaker notes and a source line on each
Every number on the slides comes from results/ (see the comments next to each); nothing is random.
Run from the repository root: python scripts/report_slides.py --template "path/to/195 Midterm Presentation.pptx"
"""
import argparse, copy, pathlib
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn
from lxml import etree

RES, FIG = pathlib.Path("results"), pathlib.Path("report/figures")
BODY, MUTED, HEAD_FILL, RULE = RGBColor(0x59, 0x59, 0x59), RGBColor(0x80, 0x80, 0x80), RGBColor(0xF3, 0xF3, 0xF3), "B7B7B7"
NO_GRID = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"          # PowerPoint's "No Style, No Grid" table style


def empty_deck(template):
    prs = Presentation(template)
    ids = prs.slides._sldIdLst
    for s in list(ids):
        prs.part.drop_rel(s.get(qn("r:id"))); ids.remove(s)
    return prs


def layout(prs, name):
    return next(l for l in prs.slide_layouts if l.name == name)


def bullets(frame, items, size):
    frame.clear()
    for i, text in enumerate(items):
        p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        p.text = text; p.space_before = Pt(0); p.space_after = Pt(5)
        for r in p.runs:
            r.font.size = Pt(size); r.font.color.rgb = BODY


def source_line(slide, text):
    box = slide.shapes.add_textbox(Inches(0.34), Inches(5.12), Inches(9.32), Inches(0.4))
    tf = box.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.text = text
    for r in p.runs:
        r.font.size = Pt(8); r.font.color.rgb = MUTED


def border_bottom(cell, color=RULE, width=12700):
    tcPr = cell._tc.get_or_add_tcPr()
    ln = etree.SubElement(tcPr, qn("a:lnB"), w=str(width), cap="flat", cmpd="sng", algn="ctr")
    fill = etree.SubElement(ln, qn("a:solidFill")); etree.SubElement(fill, qn("a:srgbClr"), val=color)
    etree.SubElement(ln, qn("a:prstDash"), val="solid")


def table(slide, rows, left, top, widths, size, bold_first_col=True):
    shape = slide.shapes.add_table(len(rows), len(rows[0]), Inches(left), Inches(top), Inches(sum(widths)), Inches(0.3 * len(rows)))
    tbl = shape.table
    tbl._tbl.tblPr.find(qn("a:tableStyleId")).text = NO_GRID
    tbl.first_row = True; tbl.horz_banding = False
    for j, w in enumerate(widths):
        tbl.columns[j].width = Inches(w)
    for i, row in enumerate(rows):
        for j, text in enumerate(row):
            c = tbl.cell(i, j); c.text = text
            c.margin_left = c.margin_right = Inches(0.06); c.margin_top = c.margin_bottom = Inches(0.03)
            c.fill.solid(); c.fill.fore_color.rgb = HEAD_FILL if i == 0 else RGBColor(0xFF, 0xFF, 0xFF)
            for p in c.text_frame.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(size); r.font.color.rgb = RGBColor(0, 0, 0) if i == 0 else BODY
                    r.font.bold = i == 0 or (j == 0 and bold_first_col)
            border_bottom(c)
    return shape


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = " ".join(text.split())


def slide_validation(prs):
    s = prs.slides.add_slide(layout(prs, "TITLE_AND_BODY"))
    s.shapes.title.text = "Technical implementation: hold out whole events"
    body = s.placeholders[1]; body.left, body.top, body.width, body.height = Inches(0.34), Inches(1.12), Inches(9.32), Inches(1.7)
    bullets(body.text_frame, [
        "Each person dates up to 22 partners, so a random split puts the same person in training and test, "
        "and the scores look better than they would for new people.",
        "So each round of 5-fold cross-validation holds out whole events; each of the 21 events is held out once.",
        "Tuning (an inner cross-validation) and data preparation (medians, log transforms, standardizing) use the "
        "training events only.",
    ], 13)
    pic = FIG / "slide_cv_design.png"                       # made by scripts/report_figures.py --slides
    s.shapes.add_picture(str(pic), Inches(1.55), Inches(2.95), width=Inches(6.9))   # 6.9 x 2.14 in, ends above the source line
    source_line(s, "Grouped cross-validation: Roberts et al. (2017), Ecography 40(8):913-929. Nested cross-validation: "
                   "Varma & Simon (2006), BMC Bioinformatics 7:91; Hastie, Tibshirani & Friedman (2009), The Elements "
                   "of Statistical Learning, 2nd ed., Section 7.10.2.")
    notes(s, """Now how we actually ran the models. The key fact about this data is that people repeat: each person
        dates up to 22 partners, so the 4,184 pairs share people. If we split pairs at random, the same person would
        be in both the training and the test set, and the model could partly recognize them. That makes the scores
        look better than they would be for a matchmaker meeting new people. So we split by event instead. The
        picture shows the five rounds: in each round the orange events are held out and scored, and the model is
        fitted on the blue ones. Every event is held out exactly once. When a model has a setting to tune, we tune
        it with a second cross-validation inside the blue events only, and all data preparation, like filling in
        missing values with medians, is learned from the training events too. Nothing about the test events leaks
        into the fit.""")


def slide_models(prs):
    s = prs.slides.add_slide(layout(prs, "TITLE_ONLY"))
    s.shapes.title.text = "How each model was fitted"
    # Choices from results/part1_models_folds.csv: no spline in any fold; QDA r = 0.9 in all 10 folds; forest
    # minimum leaf 50 in all 5 pre folds. The lasso was added after the pre-registered comparisons (PREREG.md).
    rows = [["Model", "What it chooses", "How it chooses"],
            ["Logistic regression", "a curve (spline) for age, age gap, interests, preference fit", "only if BIC improves (it never did)"],
            ["Stepwise logistic", "which predictors enter", "BIC: forward, backward, both"],
            ["LDA", "nothing to tune", "priors from the training match rate"],
            ["QDA (regularized)", "shrinkage r from 0.1 to 0.9", "best AUC on inner event folds (chose 0.9)"],
            ["Random forest", "minimum leaf size 1, 5, 20 or 50; 500 trees, √p predictors per split",
             "best AUC on inner event folds (chose 50 before the date)"],
            ["Lasso (L1 logistic)*", "penalty strength, 30 values", "best AUC on inner event folds"]]
    table(s, rows, 0.34, 1.2, [2.0, 3.8, 3.5], 12)
    box = s.shapes.add_textbox(Inches(0.34), Inches(4.15), Inches(9.32), Inches(0.9))
    tf = box.text_frame; tf.word_wrap = True
    for i, t in enumerate(["Each model is fitted twice: pre (before the date) and pre+date (adds the in-date ratings).",
                           "A yes/no call uses the training match rate (about 0.165) as the cutoff, not 0.5: only 1 pair in 6 matches.",
                           "* Added after the pre-registered comparisons."]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph(); p.text = t
        for r in p.runs:
            r.font.size = Pt(12 if i < 2 else 10); r.font.color.rgb = BODY if i < 2 else MUTED
    source_line(s, "Splines, BIC, stepwise selection and the lasso: James et al. (2023), ISLP, Sections 6.1, 6.2.2 and 7.4. "
                   "Regularized QDA: Friedman (1989), JASA 84:165-175. Random forests: Breiman (2001), Machine Learning 45:5-32.")
    notes(s, """Each model has choices to make, and every choice is made inside the training events. Logistic
        regression may bend a few continuous predictors into a curve, but only when the Bayesian information
        criterion says the curve helps, and it never did. Stepwise selection adds or removes whole predictors by the
        same criterion. QDA's shrinkage, the forest's leaf size and the lasso's penalty are tuned on the inner
        folds, again grouped by event. Because only one pair in six matches, we call a pair a match when its
        predicted probability beats the training match rate, about 0.165, rather than 0.5. One telling detail: in
        every fold QDA chose the strongest shrinkage, the forest chose the largest leaves, and no curve beat a
        straight line. Before the date, the signal is too weak to support flexible fits.""")


def slide_beyond(prs):
    s = prs.slides.add_slide(layout(prs, "TITLE_ONLY"))
    s.shapes.title.text = "Beyond the textbook: what we added, and why"
    rows = [["Method", "Why we used it", "Source"],
            ["Event-grouped cross-validation", "People repeat across pairs; a random split lets the model recognize them", "Roberts et al. 2017"],
            ["Nested cross-validation", "Tuning and scoring on the same folds makes the score optimistic", "Varma & Simon 2006"],
            ["Regularized QDA", "About 550 matches per training fold is too few to estimate a full covariance matrix, so we shrink it toward a simpler one",
             "Friedman 1989"],
            ["Lasso for logistic regression", "The textbook's lasso penalty on the logistic likelihood; drops weak predictors automatically",
             "Friedman, Hastie & Tibshirani 2010"],
            ["Standard errors clustered by event", "Pairs from one evening do well or badly together, so the usual standard error is too small",
             "Liang & Zeger 1986; Cameron & Miller 2015"],
            ["Clustered standard error of the AUC", "The AUC is the chance a random match outranks a random non-match; its standard error is also added up by event",
             "Hanley & McNeil 1982; Obuchowski 1997"],
            ["Pre-registration", "Every hypothesis and threshold committed before its test: no moving the goalposts", "Nosek et al. 2018"]]
    table(s, rows, 0.34, 1.2, [2.45, 4.75, 2.1], 10.5)
    source_line(s, "Roberts et al. (2017) Ecography 40:913-929. Varma & Simon (2006) BMC Bioinformatics 7:91. Friedman (1989) JASA "
                   "84:165-175. Friedman et al. (2010) J. Stat. Softw. 33(1). Liang & Zeger (1986) Biometrika 73:13-22. Cameron & "
                   "Miller (2015) J. Human Resources 50:317-372. Hanley & McNeil (1982) Radiology 143:29-36. Obuchowski (1997) "
                   "Biometrics 53:567-578. Nosek et al. (2018) PNAS 115:2600-2606.")
    notes(s, """These are the tools we used that go a little beyond the course textbook, each for a specific reason.
        Event-grouped and nested cross-validation are what I just described. Regularized QDA exists because QDA
        needs a separate covariance matrix for the matches, and with only about 550 matches in a training fold
        that matrix is estimated poorly, so we shrink it toward a simpler one. The lasso for logistic regression is
        the same penalty as in the textbook, applied to the logistic likelihood. Our standard errors are clustered
        by event, because pairs from the same evening tend to be predicted well or badly together, so treating them
        as independent would make our error bars too small. The AUC's standard error is clustered the same way.
        Finally, every hypothesis and its threshold was committed to our repository before we ran its test, so the
        results Jenna showed are not tuned after the fact.""")


def slide_next(prs):
    s = prs.slides.add_slide(layout(prs, "TITLE_ONLY"))
    s.shapes.title.text = "Next steps"
    rows = [["When", "What"],
            ["Weeks 6-7", "Honest error bars. Pairs share people and events, so the usual cross-validation error bars may be "
                          "too narrow. By how much, and which ones can we trust? Checked on simulated data built like ours."],
            ["Week 7", "Two follow-up questions. Who is most desirable? Rank participants within each event, with honest "
                       "uncertainty about each rank. Do people know who likes them? Score participants' own guesses that a "
                       "partner would say yes."],
            ["Week 7", "More models and checks: single trees, bagging and boosting; does the order of the evening matter (recency)?"],
            ["Week 8", "Freeze the results, robustness checks, first draft of the poster."],
            ["Weeks 9-10", "Final report, poster and presentation."]]
    table(s, rows, 0.34, 1.2, [1.4, 7.9], 12.5)
    source_line(s, "Cross-validation error bars: Bates, Hastie & Tibshirani (2024), JASA 119:1434-1445. Trees, bagging and "
                   "boosting: James et al. (2023), ISLP, Chapter 8.")
    notes(s, """Here is the plan for the second half. First, honest error bars. Because the same people appear in many
        pairs, the usual error bars on a cross-validated result may be too narrow. We want to know by how much, and
        which error bars can be trusted, and we check that on simulated data built to look like ours. Second, two
        follow-up questions: who is most desirable, ranked within each event with honest uncertainty about each
        rank, and whether people know who likes them, which we can score because participants guessed whether each
        partner would say yes. In week 7 we also add single trees, bagging and boosting, and check whether the order
        of the evening matters. Week 8 we freeze the results and draft the poster, and weeks 9 and 10 are the final
        report, poster and presentation.""")


def main():
    ap = argparse.ArgumentParser(); ap.add_argument("--template", required=True); args = ap.parse_args()
    prs = empty_deck(args.template)
    slide_validation(prs); slide_models(prs); slide_beyond(prs); slide_next(prs)
    out = RES / "report_slides.pptx"; prs.save(out); print(f"wrote {out} ({len(prs.slides)} slides)")


if __name__ == "__main__":
    main()
